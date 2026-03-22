import os
from docx import Document

resumes = [
    {
        "role": "Full Stack Developer",
        "filename": "full_stack_developer.docx",
        "power_words": ["React", "Node.js", "Python", "Microservices", "REST APIs", "Docker", "Agile", "CI/CD", "Git", "SQL"]
    },
    {
        "role": "Data Scientist",
        "filename": "data_scientist.docx",
        "power_words": ["Machine Learning", "Python", "R", "SQL", "Predictive Modeling", "Statistical Analysis", "TensorFlow", "Pandas", "NLP", "Scikit-Learn"]
    },
    {
        "role": "Data Analyst",
        "filename": "data_analyst.docx",
        "power_words": ["SQL", "Tableau", "Power BI", "Excel", "Data Visualization", "ETL", "Dashboarding", "Data Cleansing", "Python", "Statistical Analysis"]
    },
    {
        "role": "Machine Learning Engineer",
        "filename": "machine_learning_engineer.docx",
        "power_words": ["PyTorch", "TensorFlow", "Docker", "Algorithms", "Computer Vision", "NLP", "Deep Learning", "Model Evaluation", "Python", "Data Preprocessing"]
    },
    {
        "role": "Data Engineer",
        "filename": "data_engineer.docx",
        "power_words": ["Spark", "Hadoop", "Airflow", "ETL Pipelines", "Data Warehousing", "Snowflake", "Databases", "Kafka", "Python", "SQL"]
    },
    {
        "role": "Software Engineer",
        "filename": "software_engineer.docx",
        "power_words": ["Java", "C++", "Python", "Algorithms", "Data Structures", "System Design", "Agile", "Git", "RESTful Web Services", "Object-Oriented Programming"]
    },
    {
        "role": "Business Analyst",
        "filename": "business_analyst.docx",
        "power_words": ["Stakeholder Management", "Requirements Gathering", "UAT", "Process Improvement", "Jira", "Agile", "Impact Analysis", "Cross-functional Coordination", "SQL", "Data Modeling"]
    },
    {
        "role": "AI Engineer",
        "filename": "ai_engineer.docx",
        "power_words": ["Large Language Models (LLMs)", "Generative AI", "Prompt Engineering", "OpenAI API", "Hugging Face", "LangChain", "RAG", "Python", "Deep Learning", "Model Tuning"]
    },
    {
        "role": "Cloud Engineer",
        "filename": "cloud_engineer.docx",
        "power_words": ["AWS", "Azure", "GCP", "Kubernetes", "Docker", "Terraform", "Infrastructure as Code (IaC)", "CI/CD", "Linux", "Networking"]
    }
]

def generate_resumes():
    output_dir = os.path.join(os.path.dirname(__file__), 'static', 'resumes')
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    for res in resumes:
        doc = Document()
        
        # Header (Personal Details Placeholders)
        doc.add_heading('[FULL NAME]', 0)
        p = doc.add_paragraph()
        p.add_run('[City, State] | [Phone Number] | [Email Address] | [LinkedIn Profile URL] | [GitHub/Portfolio URL]\n')
        
        # Summary for Fresher
        doc.add_heading('Summary', level=1)
        doc.add_paragraph(f'Highly motivated and analytical {res["role"]} graduate with a strong academic foundation. Passionate about applying theoretical knowledge to build practical solutions and eager to contribute effectively to a dynamic, forward-thinking tech team.')
        
        # Skills
        doc.add_heading('Technical Skills', level=1)
        skills_paragraph = doc.add_paragraph()
        for word in res["power_words"]:
            skills_paragraph.add_run(f'• {word}   ')
            
        # Education First (Standard for Students)
        doc.add_heading('Education', level=1)
        doc.add_heading('[Degree] in [Major, e.g. Computer Science] | [University Name]', level=2)
        doc.add_paragraph('Expected Graduation: [Month, Year] | CGPA: [X.XX/4.00]')
        doc.add_paragraph('• Relevant Coursework: [Data Structures, Algorithms, Software Engineering, etc.]')
        
        # Projects Section (Crucial for Freshers)
        doc.add_heading('Academic & Personal Projects', level=1)
        
        doc.add_heading(f'[Key Project Name] | [Month, Year] - [Month, Year]', level=2)
        p_word1 = res["power_words"][0]
        p_word2 = res["power_words"][1]
        doc.add_paragraph(f'• Developed a comprehensive {res["role"].lower()} application leveraging {p_word1} and {p_word2} to solve [Specific Problem].')
        doc.add_paragraph(f'• Implemented core functionalities from scratch, achieving a 20% improvement in processing logic.')
        
        doc.add_heading(f'[Secondary Project Name] | [Month, Year] - [Month, Year]', level=2)
        p_word3 = res["power_words"][len(res["power_words"])-2] if len(res["power_words"]) > 2 else "advanced analytical techniques"
        p_word4 = res["power_words"][len(res["power_words"])-1] if len(res["power_words"]) > 1 else "modern tools"
        doc.add_paragraph(f'• Designed an optimized architecture utilizing {p_word3} alongside {p_word4}.')
        doc.add_paragraph(f'• Collaborated in a team of [Number] students following Agile methodologies, successfully delivering the prototype 1 week early.')
        
        # Internships / Experience (Softened for freshers)
        doc.add_heading('Internships & Leadership Experience', level=1)
        doc.add_heading(f'[Internship Title, e.g. Software Engineering Intern] | [Company/Organization Name, or Student Club]', level=2)
        doc.add_paragraph(f'• Assisted senior engineers in writing clean, maintainable code using {res["power_words"][4]} and {res["power_words"][5]}.')
        doc.add_paragraph(f'• Participated in daily stand-ups and contributed to code reviews and documentation.')
        
        # Certifications / Extracurriculars
        doc.add_heading('Extracurricular Activities & Certifications', level=1)
        doc.add_paragraph('• [Name of Certification, e.g. AWS Cloud Practitioner / DeepLearning.AI Certificate]')
        doc.add_paragraph('• Member of [University Tech Club/Society]')
        
        output_path = os.path.join(output_dir, res["filename"])
        doc.save(output_path)
        print(f'Generated fresher-friendly {output_path}')

if __name__ == "__main__":
    generate_resumes()
